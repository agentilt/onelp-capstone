#!/usr/bin/env python3
"""
Build the capstone summary slide deck (PPTX) tying the three deliverables together.

Pulls live numbers from each deliverable's results JSON and embeds the generated
figures. On-brand palette (deep slate + teal). 16:9.

Output: slides/OneLP-Capstone-Deliverables.pptx
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image as PILImage

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent

INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_LT = RGBColor(0x5E, 0xEA, 0xD4)
AMBER = RGBColor(0xD9, 0x77, 0x06)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _load():
    d1 = json.load(open(ROOT / "deliverable-1-extraction-accuracy" / "results" / "extraction_metrics.json"))
    d2 = json.load(open(ROOT / "deliverable-2-forecasting-backtest" / "results" / "forecast_metrics.json"))
    d3 = json.load(open(ROOT / "deliverable-3-pipeline-analysis" / "results" / "pipeline_analysis.json"))
    return d1, d2, d3


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide, x, y, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    return shp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4):
    """runs: list of paragraphs, each a list of (text, size, color, bold)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            r.font.name = "Helvetica Neue"
    return tb


def title_slide(prs, d1, d2, d3):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, INK)
    _box(s, Inches(0), Inches(3.05), Inches(13.333), Pt(4), fill=TEAL)
    _text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(1.4),
          [[("OneLP Capstone", 44, WHITE, True)],
           [("Proof-of-Concept Deliverables", 30, TEAL_LT, False)]])
    _text(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.6),
          [[("AI document-intelligence pipeline + portfolio analytics for Limited Partners",
             16, RGBColor(0xCB, 0xD5, 0xE1), False)]])
    _text(s, Inches(0.9), Inches(5.7), Inches(11.5), Inches(1.0),
          [[("Kishan Dhulashia   ·   Max Vanderlinden   ·   Sofia Mollon   ·   Agustin Gentil",
             14, RGBColor(0x94, 0xA3, 0xB8), False)],
           [("June 2026", 13, MUTED, False)]])


def scoreboard_slide(prs, d1, d2, d3):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    _header(s, "Three deliverables — all targets met")
    cards = [
        ("Deliverable 1", "Extraction Accuracy",
         f"{d1['summary']['hybrid']['macro_f1']:.3f}", "Hybrid macro F1",
         "Target ≥ 0.90  ·  critical-field 0.968 (≥ 0.95)"),
        ("Deliverable 2", "Forecast Backtest",
         f"{d2['backtest']['mae_pct']*100:.1f}%", "MAE vs actual",
         "Target < 15%  ·  1,000-path Monte Carlo"),
        ("Deliverable 3", "Pipeline & Latency",
         f"{d3['latency']['p50']:.0f}s", "End-to-end P50",
         "Target < 30s  ·  P95 29s  ·  ECE 0.056"),
    ]
    x = Inches(0.7)
    for tag, name, big, sub, foot in cards:
        _box(s, x, Inches(1.9), Inches(3.85), Inches(3.6), fill=LIGHT)
        _box(s, x, Inches(1.9), Inches(3.85), Pt(5), fill=TEAL)
        _text(s, x + Inches(0.3), Inches(2.2), Inches(3.3), Inches(0.8),
              [[(tag.upper(), 12, TEAL, True)], [(name, 18, INK, True)]])
        _text(s, x + Inches(0.3), Inches(3.3), Inches(3.3), Inches(1.1),
              [[(big, 48, INK, True)], [(sub, 13, MUTED, False)]])
        _text(s, x + Inches(0.3), Inches(4.75), Inches(3.3), Inches(0.7),
              [[("✓ " + foot, 11, SLATE, False)]])
        x += Inches(4.13)


def _header(slide, title, kicker=None):
    _box(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.25), fill=INK)
    _box(slide, Inches(0), Inches(1.25), Inches(13.333), Pt(3), fill=TEAL)
    runs = [[(title, 26, WHITE, True)]]
    if kicker:
        runs = [[(kicker, 12, TEAL_LT, True)], [(title, 24, WHITE, True)]]
    _text(slide, Inches(0.7), Inches(0.18), Inches(12), Inches(1.0), runs,
          anchor=MSO_ANCHOR.MIDDLE)


def problem_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    _header(s, "The problem & the system", kicker="CONTEXT")
    _text(s, Inches(0.7), Inches(1.6), Inches(5.9), Inches(5),
          [[("The problem", 18, TEAL, True)],
           [("LPs & family offices manage private-market portfolios across GP portals, "
             "PDFs, emails and spreadsheets — no standard schema.", 14, SLATE, False)],
           [("• 30–40 hours/month of manual reporting", 13, SLATE, False)],
           [("• Missed capital calls & compliance risk", 13, SLATE, False)],
           [("• No real-time exposure / liquidity view", 13, SLATE, False)]])
    _text(s, Inches(6.9), Inches(1.6), Inches(5.7), Inches(5),
          [[("The OneLP system (3 layers)", 18, TEAL, True)],
           [("1.  Ingestion — 2-stage classify, hybrid OCR+LLM extract, "
             "merge, validate, auto-apply", 13, SLATE, False)],
           [("2.  Analytics — Monte Carlo forecasting, J-curve profiles, "
             "stress tests, liquidity reserves", 13, SLATE, False)],
           [("3.  Automation — alerts, daily briefing, scheduled reports, "
             "AI email drafting", 13, SLATE, False)],
           [("This capstone evaluates the ingestion pipeline (D1, D3) and the "
             "analytics engine (D2).", 13, MUTED, False)]])


def _place_fit(slide, path, bx, by, bw, bh):
    """Place an image scaled to FIT within box (bx,by,bw,bh) [inches], preserving
    aspect ratio and centered — so it never overflows or overlaps neighbours."""
    if not Path(path).exists():
        return
    with PILImage.open(path) as im:
        ar = im.width / im.height
    box_ar = bw / bh
    if ar >= box_ar:          # width-bound
        w, h = bw, bw / ar
    else:                     # height-bound
        h, w = bh, bh * ar
    x = bx + (bw - w) / 2
    y = by + (bh - h) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def _caption(slide, text):
    _text(slide, Inches(0.7), Inches(6.95), Inches(12), Inches(0.4),
          [[(text, 11, MUTED, False)]])


def _stats_col(slide, stats, x=8.55, w=4.35):
    y = 1.95
    for label, value in stats:
        _box(slide, Inches(x), Inches(y), Inches(w), Inches(0.95), fill=LIGHT)
        _box(slide, Inches(x), Inches(y), Pt(5), Inches(0.95), fill=TEAL)
        _text(slide, Inches(x + 0.25), Inches(y + 0.08), Inches(w - 0.4), Inches(0.8),
              [[(value, 22, INK, True), ("  " + label, 12, MUTED, False)]],
              anchor=MSO_ANCHOR.MIDDLE)
        y += 1.15


def one_image_slide(prs, kicker, title, path, stats=None, caption=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    _header(s, title, kicker=kicker)
    if stats:
        _place_fit(s, path, 0.45, 1.6, 7.7, 4.95)   # left box, clear of stats column
        _stats_col(s, stats)
    else:
        _place_fit(s, path, 0.8, 1.6, 11.7, 4.95)   # full-width centered
    if caption:
        _caption(s, caption)
    return s


def two_image_slide(prs, kicker, title, p1, p2, caption=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, WHITE)
    _header(s, title, kicker=kicker)
    _place_fit(s, p1, 0.4, 1.6, 6.15, 4.95)
    _place_fit(s, p2, 6.78, 1.6, 6.15, 4.95)
    if caption:
        _caption(s, caption)
    return s


def conclusions_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, INK)
    _box(s, Inches(0), Inches(1.25), Inches(13.333), Pt(3), fill=TEAL)
    _text(s, Inches(0.7), Inches(0.3), Inches(12), Inches(0.8),
          [[("Conclusions", 28, WHITE, True)]])
    pts = [
        ("Architecture choices are evidence-backed.",
         "Hybrid extraction and Monte Carlo each beat the simpler alternative on the metric that matters."),
        ("The system is honest about uncertainty.",
         "Conservatively calibrated (ECE 0.056); classification errors are 'safe' adjacent-type confusions."),
        ("Two concrete fixes carry forward.",
         "Widen forecast intervals (65% coverage of an 80% band); regularize the correlation matrix."),
        ("Production numbers are wired.",
         "Every metric recomputes unchanged against live extractions / real cash flows — only inputs change."),
    ]
    y = 1.7
    for head, body in pts:
        _box(s, Inches(0.7), Inches(y + 0.05), Inches(0.12), Inches(0.95), fill=TEAL)
        _text(s, Inches(1.0), Inches(y), Inches(11.5), Inches(1.05),
              [[(head, 16, WHITE, True)], [(body, 13, RGBColor(0xCB, 0xD5, 0xE1), False)]])
        y += 1.15


def main():
    d1, d2, d3 = _load()
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H

    F1 = ROOT / "deliverable-1-extraction-accuracy" / "figures"
    F2 = ROOT / "deliverable-2-forecasting-backtest" / "figures"
    F3 = ROOT / "deliverable-3-pipeline-analysis" / "figures"

    title_slide(prs, d1, d2, d3)
    scoreboard_slide(prs, d1, d2, d3)
    problem_slide(prs)

    # D1
    one_image_slide(prs, "DELIVERABLE 1 · EXTRACTION ACCURACY",
                    "Hybrid extraction clears both accuracy targets",
                    F1 / "engine_comparison.png",
                    stats=[("Hybrid macro F1", f"{d1['summary']['hybrid']['macro_f1']:.3f}"),
                           ("Critical-field F1", f"{d1['summary']['hybrid']['critical_field_f1']:.3f}"),
                           ("Classification acc", f"{d1['classification']['accuracy']:.3f}")],
                    caption="Document AI → tables/forms · LLM → narrative · Hybrid → merge precedence "
                            "(merge.ts). Single engines fail in opposite directions.")
    two_image_slide(prs, "DELIVERABLE 1 · EXTRACTION ACCURACY",
                    "Accuracy by document type & classification",
                    F1 / "f1_by_document_type.png", F1 / "confusion_matrix.png",
                    caption="Hybrid F1 ≥ 0.94 for every type; classification confusions are only "
                            "between adjacent periodic-report types.")

    # D2
    two_image_slide(prs, "DELIVERABLE 2 · FORECASTING",
                    f"Monte Carlo forecast & backtest (MAE {d2['backtest']['mae_pct']*100:.1f}% < 15%)",
                    F2 / "forecast_fan_chart.png", F2 / "backtest_vs_actual.png",
                    caption="Engine is a line-for-line port of monte-carlo.ts. Rolling-origin backtest, "
                            "20 portfolio-quarter test points.")
    two_image_slide(prs, "DELIVERABLE 2 · FORECASTING",
                    "Stress testing & liquidity reserve",
                    F2 / "stress_scenarios.png", F2 / "liquidity_reserve.png",
                    caption=f"Self-funding in base case; ~${d2['reserve']['reserve_95']/1e6:.1f}M reserve "
                            "(95% confidence) needed to survive a Liquidity Crisis.")
    one_image_slide(prs, "DELIVERABLE 2 · FORECASTING",
                    "Sensitivity — what moves the forecast (±25%)",
                    F2 / "sensitivity_tornado.png",
                    stats=[("NAV level", f"±${abs(d2['sensitivity']['drivers'][0]['swing'])/1e6:.0f}M"),
                           ("Distribution rate", f"±${abs(d2['sensitivity']['drivers'][1]['swing'])/1e6:.0f}M"),
                           ("Capital-call rate", f"±${abs(d2['sensitivity']['drivers'][2]['swing'])/1e6:.0f}M")],
                    caption="NAV level and distribution pacing dominate; call timing matters ~3× less "
                            "for a mostly-deployed book.")
    one_image_slide(prs, "DELIVERABLE 2 · FORECASTING",
                    f"Cross-asset correlation widens the tails (+{d2['correlation_demo']['tail_widening_pct']*100:.0f}% σ)",
                    F2 / "correlation_effect.png",
                    caption="Cholesky correlation reduces diversification on a one-per-class demo — but "
                            "no-ops on the full 16-fund book (singular matrix, monte-carlo.ts:336-341): a documented fix.")

    # D3
    two_image_slide(prs, "DELIVERABLE 3 · PIPELINE ANALYSIS",
                    "Confidence calibration & auto-apply thresholds",
                    F3 / "reliability_diagram.png", F3 / "threshold_operating_curve.png",
                    caption="Conservatively calibrated (under-states reliability) → thresholds can be "
                            "relaxed after a calibration pass.")
    one_image_slide(prs, "DELIVERABLE 3 · PIPELINE ANALYSIS",
                    f"Latency meets target: P50 {d3['latency']['p50']:.0f}s / P95 {d3['latency']['p95']:.0f}s (Metric 3)",
                    F3 / "latency_breakdown.png",
                    caption="Targets P50 < 30s / P95 < 90s — both met. Document AI + LLM run in parallel, "
                            "so extraction cost is the slower engine, not the sum.")

    conclusions_slide(prs)

    out = HERE / "OneLP-Capstone-Deliverables.pptx"
    prs.save(out)
    n_slides = len(prs.slides._sldIdLst)
    print(f"  {out.name}: {n_slides} slides")

    # Also export a PDF copy via LibreOffice headless, if available.
    import shutil, subprocess
    soffice = (shutil.which("soffice")
               or ("/Applications/LibreOffice.app/Contents/MacOS/soffice"
                   if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists() else None))
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(out)],
                       check=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        print(f"  {out.stem}.pdf")
    print("Deck done.")


if __name__ == "__main__":
    main()
